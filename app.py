import io
import streamlit as st
from pptx import Presentation
from pptx.util import Inches, Pt


def generar_presentacion(titulo: str, subtitulo: str, puntos: list[str]) -> io.BytesIO:
    prs = Presentation()

    # Diapositiva 1: Título y Subtítulo
    slide_layout_title = prs.slide_layouts[0]
    slide1 = prs.slides.add_slide(slide_layout_title)
    slide1.shapes.title.text = titulo
    slide1.placeholders[1].text = subtitulo

    # Diapositiva 2: Lista de Puntos Clave
    slide_layout_content = prs.slide_layouts[1]
    slide2 = prs.slides.add_slide(slide_layout_content)
    slide2.shapes.title.text = "Resumen de Puntos Clave"

    text_frame = slide2.placeholders[1].text_frame
    text_frame.word_wrap = True

    for i, punto in enumerate(puntos):
        if i == 0:
            text_frame.text = punto
        else:
            p = text_frame.add_paragraph()
            p.text = punto

    # Guardar la presentación en memoria en un buffer de bytes
    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer


# Configuración de la página
st.set_page_config(
    page_title="Generador PPTX",
    page_icon="📊",
    layout="centered"
)

st.title("📊 Generador de Presentaciones PPTX")
st.write("Ingresa los detalles a continuación para construir y descargar tu archivo `.pptx`.")

# Formulario de entrada
with st.form("form_ppt"):
    titulo_input = st.text_input("Título principal", value="Reporte Ejecutivo 2026")
    subtitulo_input = st.text_input("Subtítulo / Autor", value="Presentado por el equipo de datos")
    
    puntos_input = st.text_area(
        "Puntos clave (ingresa uno por línea)",
        value="Incremento del 15% en ventas trimestrales\nLanzamiento exitoso del nuevo producto\nOptimizaciones en la cadena de suministro"
    )
    
    submit_button = st.form_submit_button("🔨 Generar PowerPoint")

# Procesamiento y Descarga
if submit_button:
    puntos_lista = [p.strip() for p in puntos_input.split("\n") if p.strip()]
    
    if not titulo_input:
        st.error("Por favor, ingresa al menos un título.")
    else:
        ppt_bytes = generar_presentacion(titulo_input, subtitulo_input, puntos_lista)
        
        st.success("¡Presentación generada correctamente!")
        
        st.download_button(
            label="📥 Descargar Presentación (.pptx)",
            data=ppt_bytes.getvalue(),
            file_name="presentacion.pptx",
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
        )
